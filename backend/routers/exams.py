from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session

from database import get_db
from models import Exam, Question, ExamQuestion
from schemas import (
    ExamCreate,
    ExamResponse,
    QuestionCreate,
    QuestionResponse,
    ExamQuestionCreate,
    ExamQuestionResponse
)
from auth import get_current_user


import io
import csv
import re

import pandas as pd
from pypdf import PdfReader
from docx import Document


router = APIRouter(
    prefix="/exams",
    tags=["Exams"]
)


@router.post("/", response_model=ExamResponse)
def create_exam(
    exam_data: ExamCreate,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user)
):

    # Only examiner can create examinations
    if current_user.role != "examiner":
        raise HTTPException(
            status_code=403,
            detail="Only examiners can create examinations"
        )

    # Validate values
    if exam_data.duration_minutes <= 0:
        raise HTTPException(
            status_code=400,
            detail="Duration must be greater than 0"
        )

    if exam_data.total_questions <= 0:
        raise HTTPException(
            status_code=400,
            detail="Total questions must be greater than 0"
        )

    if exam_data.maximum_marks <= 0:
        raise HTTPException(
            status_code=400,
            detail="Maximum marks must be greater than 0"
        )

    if exam_data.end_time <= exam_data.start_time:
        raise HTTPException(
            status_code=400,
            detail="End time must be after start time"
        )

    new_exam = Exam(
        examiner_id=current_user.id,
        exam_name=exam_data.exam_name,
        subject=exam_data.subject,
        duration_minutes=exam_data.duration_minutes,
        start_time=exam_data.start_time,
        end_time=exam_data.end_time,
        total_questions=exam_data.total_questions,
        maximum_marks=exam_data.maximum_marks
    )

    db.add(new_exam)
    db.commit()
    db.refresh(new_exam)

    return new_exam


@router.get("/my-exams", response_model=list[ExamResponse])
def get_my_exams(
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user)
):

    if current_user.role != "examiner":
        raise HTTPException(
            status_code=403,
            detail="Only examiners can view examinations"
        )

    exams = (
        db.query(Exam)
        .filter(Exam.examiner_id == current_user.id)
        .order_by(Exam.created_at.desc())
        .all()
    )

    return exams



@router.post("/questions/import")
async def import_questions(
    file: UploadFile = File(...),
    current_user=Depends(get_current_user)
):
    """
    Extract questions from PDF, DOCX, CSV or Excel files.

    The questions are returned to the frontend for review.
    They are NOT saved to the database here.
    """

    if current_user.role != "examiner":
        raise HTTPException(
            status_code=403,
            detail="Only examiners can import questions"
        )

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="Please select a file"
        )

    filename = file.filename.lower()

    allowed_extensions = (
        ".pdf",
        ".docx",
        ".csv",
        ".xlsx",
        ".xls"
    )

    if not filename.endswith(allowed_extensions):
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. "
                "Please upload PDF, Word, CSV or Excel."
            )
        )

    try:
        file_bytes = await file.read()

        if not file_bytes:
            raise HTTPException(
                status_code=400,
                detail="The uploaded file is empty."
            )

        if filename.endswith(".pdf"):
            questions = extract_from_pdf(file_bytes)

        elif filename.endswith(".docx"):
            questions = extract_from_docx(file_bytes)

        elif filename.endswith(".csv"):
            questions = extract_from_csv(file_bytes)

        elif filename.endswith((".xlsx", ".xls")):
            questions = extract_from_excel(file_bytes)

        else:
            questions = []

        return {
            "filename": file.filename,
            "count": len(questions),
            "questions": questions
        }

    except HTTPException:
        raise

    except Exception as error:
        print("Question import error:", error)

        raise HTTPException(
            status_code=400,
            detail=(
                "Unable to extract questions from this file. "
                "Please check the file format and try again."
            )
        )


def clean_text(value):
    if value is None:
        return ""

    text = str(value)

    text = text.replace("\r", " ")
    text = text.replace("\n", " ")

    return re.sub(r"\s+", " ", text).strip()


def normalize_column_name(value):
    text = clean_text(value).lower()

    text = text.replace("_", " ")
    text = text.replace("-", " ")

    return re.sub(r"\s+", " ", text).strip()


def find_column(row, possible_names):
    normalized = {
        normalize_column_name(key): key
        for key in row.keys()
    }

    for name in possible_names:
        normalized_name = normalize_column_name(name)

        if normalized_name in normalized:
            return normalized[normalized_name]

    # Partial matching
    for column in row.keys():
        normalized_column = normalize_column_name(column)

        for name in possible_names:
            normalized_name = normalize_column_name(name)

            if normalized_name in normalized_column:
                return column

    return None


def get_value(row, possible_names):
    column = find_column(
        row,
        possible_names
    )

    if column is None:
        return ""

    return clean_text(row.get(column))


def normalize_question_type(value):
    value = clean_text(value).lower()

    if value in ["mcq", "multiple choice", "multiple choice question"]:
        return "MCQ"

    if value in [
        "true false",
        "true/false",
        "true or false",
        "boolean"
    ]:
        return "True/False"

    if value in [
        "short answer",
        "short response"
    ]:
        return "Short Answer"

    if value in [
        "essay",
        "long answer",
        "long response"
    ]:
        return "Essay"

    return "MCQ"


def normalize_difficulty(value):
    value = clean_text(value).lower()

    if value == "easy":
        return "Easy"

    if value == "hard":
        return "Hard"

    return "Medium"


def normalize_marks(value):
    try:
        number = int(float(value))
        return max(number, 1)
    except:
        return 1


def build_question_from_row(row):
    question_text = get_value(
        row,
        [
            "question",
            "question text",
            "question_text",
            "question title",
            "title",
            "text"
        ]
    )

    if not question_text:
        return None

    option_a = get_value(
        row,
        [
            "option a",
            "option_a",
            "choice a",
            "answer a",
            "a"
        ]
    )

    option_b = get_value(
        row,
        [
            "option b",
            "option_b",
            "choice b",
            "answer b",
            "b"
        ]
    )

    option_c = get_value(
        row,
        [
            "option c",
            "option_c",
            "choice c",
            "answer c",
            "c"
        ]
    )

    option_d = get_value(
        row,
        [
            "option d",
            "option_d",
            "choice d",
            "answer d",
            "d"
        ]
    )

    correct_answer = get_value(
        row,
        [
            "correct answer",
            "correct_answer",
            "correct option",
            "answer",
            "correct",
            "right answer"
        ]
    )

    subject = get_value(
        row,
        [
            "subject",
            "course",
            "topic",
            "category"
        ]
    )

    question_type = get_value(
        row,
        [
            "question type",
            "question_type",
            "type"
        ]
    )

    difficulty = get_value(
        row,
        [
            "difficulty",
            "level"
        ]
    )

    marks = get_value(
        row,
        [
            "marks",
            "mark",
            "points",
            "score"
        ]
    )

    return {
        "subject": subject or "General",
        "question_text": question_text,
        "question_type": normalize_question_type(
            question_type
        ),
        "difficulty": normalize_difficulty(
            difficulty
        ),
        "option_a": option_a or None,
        "option_b": option_b or None,
        "option_c": option_c or None,
        "option_d": option_d or None,
        "correct_answer": correct_answer or "",
        "marks": normalize_marks(marks)
    }


def extract_from_csv(file_bytes):
    text = file_bytes.decode(
        "utf-8-sig",
        errors="replace"
    )

    sample = text[:5000]

    try:
        dialect = csv.Sniffer().sniff(sample)
    except:
        dialect = csv.excel

    reader = csv.DictReader(
        io.StringIO(text),
        dialect=dialect
    )

    questions = []

    for row in reader:
        question = build_question_from_row(row)

        if question:
            questions.append(question)

    return questions


def extract_from_excel(file_bytes):
    dataframe = pd.read_excel(
        io.BytesIO(file_bytes)
    )

    dataframe = dataframe.fillna("")

    questions = []

    for _, row in dataframe.iterrows():
        row_data = row.to_dict()

        question = build_question_from_row(
            row_data
        )

        if question:
            questions.append(question)

    return questions


def extract_from_docx(file_bytes):
    document = Document(
        io.BytesIO(file_bytes)
    )

    questions = []

    # First try tables.
    for table in document.tables:

        headers = []

        if table.rows:
            headers = [
                clean_text(cell.text)
                for cell in table.rows[0].cells
            ]

        if len(headers) >= 2:

            for table_row in table.rows[1:]:

                values = [
                    clean_text(cell.text)
                    for cell in table_row.cells
                ]

                if not any(values):
                    continue

                row = {}

                for index, header in enumerate(headers):
                    if index < len(values):
                        row[header] = values[index]

                question = build_question_from_row(row)

                if question:
                    questions.append(question)

    # If no table questions were found,
    # try normal paragraph format.
    if not questions:

        paragraphs = [
            clean_text(paragraph.text)
            for paragraph in document.paragraphs
            if clean_text(paragraph.text)
        ]

        questions = extract_questions_from_text(
            paragraphs
        )

    return questions


def extract_from_pdf(file_bytes):
    reader = PdfReader(
        io.BytesIO(file_bytes)
    )

    text_parts = []

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text_parts.append(page_text)

    full_text = "\n".join(text_parts)

    if not full_text.strip():
        raise HTTPException(
            status_code=400,
            detail=(
                "No readable text was found in the PDF. "
                "Scanned image PDFs are not supported yet."
            )
        )

    paragraphs = [
        clean_text(line)
        for line in full_text.splitlines()
        if clean_text(line)
    ]

    return extract_questions_from_text(
        paragraphs
    )


def extract_questions_from_text(lines):
    questions = []

    current = None

    question_pattern = re.compile(
        r"^(?:question\s*)?(\d+)[\.\):\-]\s*(.*)$",
        re.IGNORECASE
    )

    option_pattern = re.compile(
        r"^([A-Da-d])[\.\):\-]\s*(.*)$"
    )

    for line in lines:

        question_match = question_pattern.match(
            line
        )

        if question_match:

            if current and current.get(
                "question_text"
            ):
                questions.append(current)

            current = {
                "subject": "General",
                "question_text":
                    question_match.group(2),
                "question_type": "MCQ",
                "difficulty": "Medium",
                "option_a": None,
                "option_b": None,
                "option_c": None,
                "option_d": None,
                "correct_answer": "",
                "marks": 1
            }

            continue

        option_match = option_pattern.match(
            line
        )

        if option_match and current:

            letter = option_match.group(1).upper()
            value = option_match.group(2).strip()

            if letter == "A":
                current["option_a"] = value

            elif letter == "B":
                current["option_b"] = value

            elif letter == "C":
                current["option_c"] = value

            elif letter == "D":
                current["option_d"] = value

            continue

        if current:

            # Look for answer information.
            answer_match = re.match(
                r"^(?:correct answer|answer|correct)\s*[:\-]\s*(.*)$",
                line,
                re.IGNORECASE
            )

            if answer_match:
                current["correct_answer"] = (
                    answer_match.group(1).strip()
                )
                continue

            # Continue question text if it is not
            # an option or answer line.
            current["question_text"] += (
                " " + line
            )

    if current and current.get("question_text"):
        questions.append(current)

    return questions


@router.post(
    "/questions",
    response_model=QuestionResponse
)
def create_question(
    question_data: QuestionCreate,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user)
):

    # Only examiner can create questions
    if current_user.role != "examiner":
        raise HTTPException(
            status_code=403,
            detail="Only examiners can create questions"
        )

    # Validate marks
    if question_data.marks <= 0:
        raise HTTPException(
            status_code=400,
            detail="Marks must be greater than 0"
        )

    # Validate question text
    if not question_data.question_text.strip():
        raise HTTPException(
            status_code=400,
            detail="Question text is required"
        )

    # Validate subject
    if not question_data.subject.strip():
        raise HTTPException(
            status_code=400,
            detail="Subject is required"
        )

    # MCQ validation
    if question_data.question_type == "MCQ":

        if not all([
            question_data.option_a,
            question_data.option_b,
            question_data.option_c,
            question_data.option_d
        ]):
            raise HTTPException(
                status_code=400,
                detail="All four options are required for MCQ"
            )

    new_question = Question(
        examiner_id=current_user.id,
        subject=question_data.subject,
        question_text=question_data.question_text,
        question_type=question_data.question_type,
        difficulty=question_data.difficulty,
        option_a=question_data.option_a,
        option_b=question_data.option_b,
        option_c=question_data.option_c,
        option_d=question_data.option_d,
        correct_answer=question_data.correct_answer,
        marks=question_data.marks
    )

    db.add(new_question)
    db.commit()
    db.refresh(new_question)

    return new_question


@router.get(
    "/questions/my-questions",
    response_model=list[QuestionResponse]
)
def get_my_questions(
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user)
):

    if current_user.role != "examiner":
        raise HTTPException(
            status_code=403,
            detail="Only examiners can view questions"
        )

    questions = (
        db.query(Question)
        .filter(
            Question.examiner_id == current_user.id
        )
        .order_by(
            Question.created_at.desc()
        )
        .all()
    )

    return questions


@router.post(
    "/{exam_id}/questions",
    response_model=list[ExamQuestionResponse]
)
def add_questions_to_exam(
    exam_id: int,
    question_data: ExamQuestionCreate,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user)
):

    if current_user.role != "examiner":
        raise HTTPException(
            status_code=403,
            detail="Only examiners can add questions"
        )

    # Check that the exam belongs to this examiner
    exam = (
        db.query(Exam)
        .filter(
            Exam.id == exam_id,
            Exam.examiner_id == current_user.id
        )
        .first()
    )

    if not exam:
        raise HTTPException(
            status_code=404,
            detail="Examination not found"
        )

    if not question_data.question_ids:
        raise HTTPException(
            status_code=400,
            detail="Please select at least one question"
        )

    # Prevent duplicate question IDs
    unique_question_ids = list(
        dict.fromkeys(question_data.question_ids)
    )

    # Check that all questions belong to this examiner
    questions = (
        db.query(Question)
        .filter(
            Question.id.in_(unique_question_ids),
            Question.examiner_id == current_user.id
        )
        .all()
    )

    if len(questions) != len(unique_question_ids):
        raise HTTPException(
            status_code=400,
            detail="One or more questions are invalid"
        )

    # Remove previous question assignments
    db.query(ExamQuestion).filter(
        ExamQuestion.exam_id == exam_id
    ).delete(
        synchronize_session=False
    )

    # Randomize question order
    import random

    randomized_ids = unique_question_ids.copy()
    random.shuffle(randomized_ids)

    exam_questions = []

    for order, question_id in enumerate(
        randomized_ids,
        start=1
    ):
        exam_question = ExamQuestion(
            exam_id=exam_id,
            question_id=question_id,
            question_order=order
        )

        db.add(exam_question)
        exam_questions.append(exam_question)

    db.commit()

    for item in exam_questions:
        db.refresh(item)

    return exam_questions


@router.get(
    "/{exam_id}/questions",
    response_model=list[ExamQuestionResponse]
)
def get_exam_questions(
    exam_id: int,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user)
):
    if current_user.role != "examiner":
        raise HTTPException(
            status_code=403,
            detail="Only examiners can view examination questions"
        )

    exam = (
        db.query(Exam)
        .filter(
            Exam.id == exam_id,
            Exam.examiner_id == current_user.id
        )
        .first()
    )

    if not exam:
        raise HTTPException(
            status_code=404,
            detail="Examination not found"
        )

    exam_questions = (
        db.query(ExamQuestion)
        .filter(
            ExamQuestion.exam_id == exam_id
        )
        .order_by(
            ExamQuestion.question_order.asc()
        )
        .all()
    )

    return exam_questions